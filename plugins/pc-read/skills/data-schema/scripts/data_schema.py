#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.12"
# dependencies = ["openpyxl", "pyarrow"]
# ///
"""Print the schema and the first rows of a data file: ``.xlsx``, ``.csv`` or ``.parquet``.

For each sheet (a ``.csv`` or ``.parquet`` is one sheet) the output lists the columns with their
type, then at most ``SAMPLE_ROWS`` rows. A workbook weighs by its cell count, not by its format;
what a model reads usefully is the shape, a sample, then targeted queries.

The sample never goes through ``Read``, so the ``pc-privacy`` guard cannot see it. This script asks
that guard first (``pii_guard.py --verdict``): a flagged file, or a guard that fails, times out or
prints something else than JSON, gives the schema alone. Only ``pc-privacy`` being absent, or a
successful ``flagged: false``, gives the sample.
"""

from __future__ import annotations

import argparse
import csv
import datetime
import json
import os
import subprocess
import sys
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path

SAMPLE_ROWS = 5
TYPE_WINDOW = 100  # rows looked at to name a column's type
CELL_WIDTH = 60
GUARD_TIMEOUT_S = 55  # the guard loads Presidio + spaCy cold
SUFFIXES = {".xlsx", ".csv", ".parquet"}

Sheet = tuple[
    str, list[str], list[str], list[list[str]]
]  # name, columns, types, sample rows


def _type_of(value: object) -> str | None:
    """Name the type of one cell.

    Args:
        value: A cell value, ``None`` or an empty string meaning no value.

    Returns:
        A type name, or None for an empty cell.
    """
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, int):
        return "int"
    if isinstance(value, float):
        return "float"
    if isinstance(value, datetime.datetime):
        return "datetime"
    if isinstance(value, datetime.date):
        return "date"
    if isinstance(value, str):
        for kind, cast in (("int", int), ("float", float)):
            try:
                cast(value)
                return kind
            except ValueError:
                pass
    return "str"


def _merge(seen: set[str]) -> str:
    """Collapse the types seen in a column into one name.

    Args:
        seen: Type names found in the column.

    Returns:
        ``empty`` when none, the type when one, ``float`` for int + float, else ``mixed``.
    """
    if not seen:
        return "empty"
    if len(seen) == 1:
        return next(iter(seen))
    if seen == {"int", "float"}:
        return "float"
    return "mixed"


def _cell(value: object) -> str:
    """Render a cell on one line, bounded, safe inside a Markdown table.

    Args:
        value: A cell value.

    Returns:
        The text; empty for no value.
    """
    if value is None:
        return ""
    text = " ".join(str(value).split()).replace("|", "\\|")
    return text if len(text) <= CELL_WIDTH else text[: CELL_WIDTH - 1] + "…"


def _shape(
    name: str, rows: Iterator[list[object]], header: list[str] | None = None
) -> Sheet:
    """Build a sheet from its rows: header, types over a window, a short sample.

    Args:
        name: The sheet name.
        rows: The rows in order; only the first ones are consumed.
        header: Column names when the source carries them apart from the rows.

    Returns:
        The sheet; fully empty columns are dropped.
    """
    cells: list[list[object]] = []
    for row in rows:
        if not any(v is not None and v != "" for v in row):
            continue
        if header is None:
            header = [_cell(v) for v in row]
            continue
        cells.append(list(row))
        if len(cells) >= TYPE_WINDOW:
            break
    names = header or []
    width = max([len(names), *(len(r) for r in cells)])
    names = names + [""] * (width - len(names))
    keep: list[int] = []
    types: dict[int, str] = {}
    for i in range(width):
        seen = {t for r in cells if i < len(r) and (t := _type_of(r[i]))}
        if names[i] or seen:
            keep.append(i)
            types[i] = _merge(seen)
    columns = [names[i] or f"column_{i + 1}" for i in keep]
    sample = [
        [_cell(r[i]) if i < len(r) else "" for i in keep] for r in cells[:SAMPLE_ROWS]
    ]
    return name, columns, [types[i] for i in keep], sample


def read_xlsx(path: Path) -> list[Sheet]:
    """Read every sheet of a workbook, lazily.

    Args:
        path: The ``.xlsx`` file.

    Returns:
        One sheet per worksheet, in workbook order.
    """
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return [
            _shape(ws.title, (list(r) for r in ws.iter_rows(values_only=True)))
            for ws in book.worksheets
        ]
    finally:
        book.close()


def read_csv(path: Path) -> list[Sheet]:
    """Read a CSV file as one sheet.

    Args:
        path: The ``.csv`` file.

    Returns:
        A single sheet named after the file.
    """
    with path.open(newline="", encoding="utf-8-sig", errors="replace") as handle:
        head = handle.read(8192)
        handle.seek(0)
        try:
            dialect: type[csv.Dialect] = csv.Sniffer().sniff(head, delimiters=",;\t|")  # type: ignore[assignment]
        except csv.Error:
            dialect = csv.excel
        return [_shape(path.stem, (list(r) for r in csv.reader(handle, dialect)))]


def read_parquet(path: Path) -> list[Sheet]:
    """Read a Parquet file as one sheet; its types come from the file's schema.

    Args:
        path: The ``.parquet`` file.

    Returns:
        A single sheet named after the file.
    """
    import pyarrow.parquet as pq

    file = pq.ParquetFile(path)
    columns = [f.name for f in file.schema_arrow]
    types = [str(f.type) for f in file.schema_arrow]
    sample: list[list[str]] = []
    for batch in file.iter_batches(batch_size=SAMPLE_ROWS):
        for row in batch.to_pylist():
            sample.append([_cell(row[c]) for c in columns])
        break
    return [(path.stem, columns, types, sample[:SAMPLE_ROWS])]


READERS = {".xlsx": read_xlsx, ".csv": read_csv, ".parquet": read_parquet}


def privacy_root(explicit: str | None) -> str | None:
    """Find the folder of the ``pc-privacy`` plugin, if installed.

    Args:
        explicit: A folder given on the command line (the one ``pc-privacy`` published).

    Returns:
        The folder, or None when ``pc-privacy`` is not installed.
    """
    if explicit:
        return explicit
    config = Path(os.environ.get("CLAUDE_CONFIG_DIR") or "~/.claude").expanduser()
    try:
        plugins = json.loads(
            (config / "plugins" / "installed_plugins.json").read_text()
        )["plugins"]
    except (OSError, ValueError, KeyError, TypeError):
        return None
    try:
        known = json.loads((config / "plugins" / "known_marketplaces.json").read_text())
    except (OSError, ValueError):
        known = {}
    for key, installs in plugins.items():
        if not key.startswith("pc-privacy@"):
            continue
        entry = known.get(key.split("@", 1)[1]) if isinstance(known, dict) else None
        source = entry.get("source") if isinstance(entry, dict) else None
        tree = (
            source.get("path")
            if isinstance(source, dict) and source.get("source") == "directory"
            else None
        )
        if isinstance(tree, str):
            candidate = str(Path(tree) / "plugins" / "pc-privacy")
            if (Path(candidate) / "hooks" / "pii_guard.py").is_file():
                return candidate
        for install in installs if isinstance(installs, list) else []:
            root = install.get("installPath") if isinstance(install, dict) else None
            if (
                isinstance(root, str)
                and (Path(root) / "hooks" / "pii_guard.py").is_file()
            ):
                return root
    return None


@dataclass(frozen=True)
class SampleVerdict:
    """Whether sample rows may be shown, and why."""

    allowed: bool
    reason: str


def sample_allowed(path: Path, root: str | None) -> SampleVerdict:
    """Ask the ``pc-privacy`` guard whether rows of ``path`` may be shown.

    Args:
        path: The data file.
        root: The ``pc-privacy`` folder, or None when it is absent.

    Returns:
        The verdict, with its reason in a few words. Only an absent guard or a
        successful ``flagged: false`` allows it.
    """
    if root is None:
        return SampleVerdict(True, "pc-privacy not installed")
    event = {"tool_name": "Read", "tool_input": {"file_path": str(path)}}
    try:
        done = subprocess.run(
            ["uv", "run", "--script", f"{root}/hooks/pii_guard.py", "--verdict"],
            input=json.dumps(event),
            capture_output=True,
            text=True,
            timeout=GUARD_TIMEOUT_S,
            check=False,
        )
        if done.returncode == 0 and json.loads(done.stdout).get("flagged") is False:
            return SampleVerdict(True, "pc-privacy found no personal data")
        if done.returncode == 0 and json.loads(done.stdout).get("flagged") is True:
            return SampleVerdict(False, "pc-privacy flags this file as personal data")
        return SampleVerdict(False, "pc-privacy guard gave no usable answer")
    except (OSError, ValueError, AttributeError, subprocess.TimeoutExpired):
        return SampleVerdict(False, "pc-privacy guard failed or timed out")


def render(sheets: list[Sheet], show_rows: bool, note: str) -> str:
    """Render the sheets as Markdown.

    Args:
        sheets: The sheets to render.
        show_rows: Whether to include sample rows.
        note: Why rows are withheld, shown when they are.

    Returns:
        The Markdown text.
    """
    out: list[str] = []
    for name, columns, types, sample in sheets:
        out.append(f"## Sheet: {name}")
        out.append("")
        out.append("| column | type |")
        out.append("|---|---|")
        out.extend(f"| {_cell(c)} | {t} |" for c, t in zip(columns, types, strict=True))
        out.append("")
        if show_rows and columns:
            out.append(f"First {len(sample)} row(s):")
            out.append("")
            out.append("| " + " | ".join(_cell(c) for c in columns) + " |")
            out.append("|" + "---|" * len(columns))
            out.extend("| " + " | ".join(r) + " |" for r in sample)
            out.append("")
        elif not show_rows:
            out.append(f"Rows withheld: {note}. Schema only.")
            out.append("")
    return "\n".join(out)


def main(argv: list[str]) -> int:
    """Print the schema and sample of a data file.

    Args:
        argv: The arguments after the script name.

    Returns:
        0 on success, 2 on an unusable file.
    """
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--privacy-root", help="folder of the pc-privacy plugin")
    parser.add_argument("--out", help="write the Markdown here instead of stdout")
    parser.add_argument("path")
    args = parser.parse_args(argv)
    path = Path(args.path)
    if path.suffix.lower() not in SUFFIXES or not path.is_file():
        print(
            f"data_schema: not a readable {', '.join(sorted(SUFFIXES))} file: {path}",
            file=sys.stderr,
        )
        return 2
    try:
        sheets = READERS[path.suffix.lower()](path)
    except Exception as exc:  # a corrupt file is the caller's answer, not a traceback
        print(
            f"data_schema: cannot read {path}: {type(exc).__name__}: {exc}",
            file=sys.stderr,
        )
        return 2
    verdict = sample_allowed(path, privacy_root(args.privacy_root))
    text = render(sheets, verdict.allowed, verdict.reason)
    if args.out:
        Path(args.out).write_text(text, encoding="utf-8")
        print(args.out)
    else:
        print(text)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
